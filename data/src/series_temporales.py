"""
series_temporales.py
=====================
La hoja "Historico niveles" de `base_final.xlsx` es una tabla ancha: una
fila por pozo (Cod UK) y una columna por fecha de medición. Los
encabezados de esas columnas de fecha están escritos de forma muy
inconsistente (texto libre tipo "Oct-86", "Dic95", "sept-14", "Nov_16",
"Mar.2017"... y, para otro tramo, fechas de Excel donde el "día" del
valor en realidad representa el año de 2 dígitos, p. ej. la celda se ve
como "Jun-10" pero el año real es 2010, no el año calendario que Excel
le puso por defecto al guardar la fecha).

Este módulo decodifica esos encabezados a una fecha real (año-mes) y arma
una tabla larga (Cod UK, fecha, nivel) lista para graficar series
temporales por pozo. Requiere `openpyxl` para leer el `number_format` de
cada celda de encabezado, que es la pista que permite distinguir ambos
casos.
"""
from __future__ import annotations

import datetime as _dt
import re

import pandas as pd

from data_loading import BASE_FINAL_PATH

_MESES = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "sep": 9, "oct": 10, "nov": 11, "dic": 12,
}
_TEXTO_RE = re.compile(r"([A-Za-zé]+)[\.\-_ ]*'?(\d{2,4})")

# Columnas de la hoja que NO son fechas de medición.
_COLUMNAS_NO_FECHA = {"cod uk", "z_m"}


def _mes_desde_texto(texto: str) -> int | None:
    return _MESES.get(texto.lower()[:3])


def _parsear_fecha_texto(valor: str) -> pd.Timestamp | None:
    match = _TEXTO_RE.search(valor.strip())
    if not match:
        return None
    mes = _mes_desde_texto(match.group(1))
    if mes is None:
        return None
    anio_str = match.group(2)
    anio = int(anio_str)
    if len(anio_str) == 2:
        anio += 2000 if anio < 50 else 1900
    try:
        return pd.Timestamp(year=anio, month=mes, day=1)
    except ValueError:
        return None


def _decodificar_fecha_columna(valor, formato: str) -> pd.Timestamp | None:
    """
    `valor` es lo que devuelve openpyxl para la celda de encabezado;
    `formato` es su `number_format`. Si el formato no muestra año
    ("mmm-d", "mmmdd", etc.), el "día" de la fecha guardada es en
    realidad el año de 2 dígitos real de la medición.
    """
    if valor is None:
        return None
    if isinstance(valor, (_dt.datetime, _dt.date)):
        formato_low = (formato or "").lower()
        if "yy" in formato_low:
            # El formato sí muestra año -> la fecha guardada es la real.
            return pd.Timestamp(year=valor.year, month=valor.month, day=1)
        # El formato solo muestra mes+día -> el "día" codifica el año (2 dígitos).
        anio_corto = valor.day
        anio = 2000 + anio_corto if anio_corto < 50 else 1900 + anio_corto
        return pd.Timestamp(year=anio, month=valor.month, day=1)
    if isinstance(valor, str):
        return _parsear_fecha_texto(valor)
    return None


def construir_serie_temporal_niveles(ruta=BASE_FINAL_PATH, hoja: str = "Historico niveles") -> pd.DataFrame:
    """
    Devuelve una tabla larga con columnas: Cod UK, fecha (Timestamp,
    normalizada al día 1 del mes), nivel (float). Se descartan las
    columnas que no se pudieron decodificar como fecha y las celdas
    vacías.
    """
    import openpyxl

    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb[hoja]

    fechas_por_columna: dict[int, pd.Timestamp] = {}
    for col in range(1, ws.max_column + 1):
        celda = ws.cell(row=1, column=col)
        etiqueta = str(celda.value).strip().lower() if celda.value is not None else ""
        if etiqueta in _COLUMNAS_NO_FECHA or etiqueta.startswith("fuente"):
            continue
        fecha = _decodificar_fecha_columna(celda.value, celda.number_format)
        if fecha is not None:
            fechas_por_columna[col] = fecha

    filas = []
    for fila in range(2, ws.max_row + 1):
        cod_uk = ws.cell(row=fila, column=1).value
        if not cod_uk:
            continue
        for col, fecha in fechas_por_columna.items():
            valor = ws.cell(row=fila, column=col).value
            if valor is None or not isinstance(valor, (int, float)):
                continue
            filas.append({"Cod UK": str(cod_uk).strip(), "fecha": fecha, "nivel": float(valor)})

    return pd.DataFrame(filas, columns=["Cod UK", "fecha", "nivel"]).sort_values(["Cod UK", "fecha"])
